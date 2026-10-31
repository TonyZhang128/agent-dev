"""离线 OCR 资料，全部由测试生成，不读取开发者的数据池。"""

import json
from pathlib import Path

import openpyxl
import pymupdf
from openpyxl.styles import Border, Side

from app.config import Settings

CODE = "TEST"


def make_ocr_settings(root: Path, *, name: str = "测试表", height: int = 300) -> Settings:
    templates, assets, projects = root / "templates", root / "assets", root / "projects"
    projects.mkdir(parents=True)
    (assets / "pdf").mkdir(parents=True)
    (assets / "excel").mkdir()
    with pymupdf.open() as document:
        page = document.new_page(width=400, height=height)
        for x in (40, 140, 240, 340):
            page.draw_line((x, 40), (x, 240))
        for y in (40, 90, 140, 190, 240):
            page.draw_line((40, y), (340, y))
        page.insert_text((45, 65), "Item")
        document.save(assets / "pdf/test.pdf")
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    line = Side(style="thin")
    for row in worksheet.iter_rows(min_row=1, max_row=5, min_col=1, max_col=3):
        for cell in row:
            cell.border = Border(top=line, bottom=line, left=line, right=line)
    worksheet.merge_cells("A1:B1")
    worksheet["A1"] = "监理单位：        编 号:TEST"
    worksheet["A2"] = "          土方路基"
    workbook.save(assets / "excel/test.xlsx")
    workbook.close()
    documents = {
        "forms/files.json": {
            "forms": {
                CODE: {"name": name, "pdf": ["pdf/test.pdf"], "excel": {"path": "excel/test.xlsx"}}
            }
        },
        "layouts/jb12.json": {
            "forms": {
                CODE: {
                    "header": [{"field": "fld_supervisor", "rect": [0.36, 0.14, 0.2, 0.14]}],
                    "key": [],
                }
            }
        },
        "layouts/book1.json": {"forms": {}},
        "layouts/jianbiao.json": {
            "forms": {
                "SIMPLE": {
                    "target": CODE,
                    "name": "测试简表",
                    "title": "测试简表",
                    "aspect": 1,
                    "cells": [
                        {
                            "field": "fld_1",
                            "rect": [0.1, 0.1, 0.8, 0.8],
                            "label": "数值",
                            "type": "number",
                        }
                    ],
                }
            }
        },
        "refs/wbs_rules.json": {
            "drawingCues": {"bridgeName": "(.+桥)", "tunnelName": "(.+隧道)", "girder": "T梁"},
            "bridge": {
                "abutment": ["基础"],
                "substructure": ["墩身"],
                "superPrecast": ["安装"],
                "superCast": ["浇筑"],
                "deck": ["桥面"],
                "protect": ["防护"],
            },
            "road": {
                "partKm": 3,
                "earthwork": ["土方路基"],
                "drainage": ["排水沟"],
                "protection": ["挡土墙"],
                "pavement": ["基层"],
            },
            "units": {"LJ": {"splitKm": 10}, "LM": {"splitKm": 10}},
        },
    }
    for relative, data in documents.items():
        path = templates / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    return Settings(
        _env_file=None,
        templates_root=templates,
        asset_root=assets,
        admin_root=root / "admin",
        projects_root=projects,
        ocr_model_root=root / "models",
    )


def fill_config() -> dict:
    fields = [
        ("fld_supervisor", "B1", [0.36, 0.14, 0.2, 0.14]),
        ("fld_location", "A2", [0.36, 0.31, 0.2, 0.14]),
        ("fld_1", "B3", [0.36, 0.48, 0.2, 0.14]),
        ("fld_2", "C3", [0.61, 0.48, 0.2, 0.14]),
        ("fld_rate", "B4", [0.36, 0.65, 0.2, 0.14]),
    ]
    anchors = []
    for field, cell, (x, y, w, h) in fields:
        anchors.extend(
            [
                {
                    "id": field,
                    "doc": CODE,
                    "kind": "PDF_REGION",
                    "field": field,
                    "rect": {"x": x, "y": y, "w": w, "h": h},
                },
                {
                    "id": field + "_excel",
                    "doc": CODE,
                    "kind": "EXCEL_RANGE",
                    "field": field,
                    "range": cell,
                },
            ]
        )
    return {
        "anchors": anchors,
        "formulas": [{"target": "fld_rate", "expr": "PASS_RATE(94, @{fld_1}, @{fld_2})"}],
    }
